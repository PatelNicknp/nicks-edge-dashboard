"""Builds the two starter workbooks from the latest output/ files.
1. Nicks_Edge_Dashboard_starter.xlsx  - import into Google Sheets; every data tab pulls live from GitHub with IMPORTDATA
2. today_snapshot.xlsx                - the same tabs filled with today's numbers (preview, or a backup)"""
import os, json, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
HERE = os.path.dirname(os.path.abspath(__file__)); OUT = os.path.join(os.path.dirname(HERE), "output")
TABS = ["top_setups", "scorecard_details", "news_risk", "upcoming_events", "econ_data", "econ_strength", "cot_report",
        "cot_history", "retail_sentiment", "put_call", "seasonality", "momentum", "volatility", "market_heatmap", "carry",
        "rates", "spreads", "calendar", "news", "score_history", "backtest", "updates_log", "help", "settings"]
DARK, PANEL, TEXT, ACCENT = "0E1117", "161B22", "E6EDF3", "2F81F7"
H = Font(bold=True, color=TEXT); FILL = PatternFill("solid", fgColor=PANEL)

def setup_sheet(wb, user=""):
    ws = wb.active; ws.title = "Setup"
    rows = [("Nick's Edge Dashboard", ""), ("Your GitHub username", user), ("Repository name", "nicks-edge-dashboard"),
            ("Branch", "main"), ("Data address (do not change)", '="https://raw.githubusercontent.com/"&B2&"/"&B3&"/"&B4&"/output/"'),
            ("Refresh helper (leave blank)", ""), ("", ""),
            ("What to do", "Type your GitHub username in B2. Every other tab fills by itself within a minute."),
            ("Updates", "The data files update every hour. Google Sheets re-reads them about once an hour."),
            ("Note", "Research screen, not a trade signal. All times are Pacific Time (PT).")]
    for i, (a, b) in enumerate(rows, 1):
        ws.cell(i, 1, a).font = H; ws.cell(i, 2, b)
    ws["A1"].font = Font(bold=True, size=16, color=ACCENT); ws["B2"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["B2"].font = Font(color="0000FF", bold=True)
    ws.column_dimensions["A"].width = 32; ws.column_dimensions["B"].width = 90
    return ws

def starter():
    wb = Workbook(); setup_sheet(wb)
    for t in TABS:
        ws = wb.create_sheet(t)
        ws["A1"] = '=IMPORTDATA(Setup!$B$5&"%s.csv"&Setup!$B$6)' % t
    p = os.path.join(HERE, "Nicks_Edge_Dashboard_starter.xlsx"); wb.save(p); return p

def snapshot():
    wb = Workbook(); setup_sheet(wb)
    meta = json.load(open(os.path.join(OUT, "meta.json")))
    wb["Setup"]["B6"] = ""; wb["Setup"]["B1"] = "Snapshot taken " + meta["updated_pt"]
    for t in TABS:
        df = pd.read_csv(os.path.join(OUT, t + ".csv"))
        ws = wb.create_sheet(t)
        ws.append(list(df.columns))
        for c in ws[1]: c.font = H; c.fill = FILL; c.alignment = Alignment(horizontal="center")
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
        ws.freeze_panes = "A2"
        for i, col in enumerate(df.columns, 1):
            w = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(w + 2, 8), 60)
    p = os.path.join(HERE, "today_snapshot.xlsx"); wb.save(p); return p

if __name__ == "__main__":
    print(starter()); print(snapshot())
